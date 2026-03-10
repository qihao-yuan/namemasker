"""
Process Word (.docx) files: NER-style type matching via LLM on raw text.
No image conversion - just extract text, detect entities, replace.
"""
import re
import json
from docx import Document


def _replace_in_runs(runs, entity: str, mask: str):
    """
    Replace entity text across potentially split runs.
    Word may split one visible string into multiple runs with different formatting.
    """
    if not runs:
        return 0

    full = "".join(r.text for r in runs)
    if entity not in full:
        return 0

    count = 0
    while entity in full:
        idx = full.index(entity)
        end = idx + len(entity)

        char_pos = 0
        for run in runs:
            run_start = char_pos
            run_end = char_pos + len(run.text)

            overlap_start = max(idx, run_start)
            overlap_end = min(end, run_end)

            if overlap_start < overlap_end:
                local_start = overlap_start - run_start
                local_end = overlap_end - run_start
                replacement = mask if overlap_start == idx else ""
                run.text = run.text[:local_start] + replacement + run.text[local_end:]

            char_pos = run_end

        full = "".join(r.text for r in runs)
        count += 1

    return count


class DocxProcessor:
    """
    Two-phase Word processor:
    Phase 1 (detect): extract text -> LLM NER -> collect entities
    Phase 2 (save):   replace entities in runs and save
    """

    def __init__(self, input_path, api_key, model, base_url, target="",
                 passes=2, mode="fill", pad_pct=0.0):
        self.input_path = input_path
        self.api_key = api_key
        self.model = model
        self.base_url = base_url
        self.target = target

        self.doc = None
        self.found_entities = []  # list of {"text": "...", "type": "..."}

    def detect(self, on_log=None, on_names=None, on_preview=None, is_cancelled=None):
        from namemasker.api_detect import DEFAULT_TARGET, DEFAULT_BASE_URL
        from openai import OpenAI

        target = self.target or DEFAULT_TARGET
        base_url = self.base_url or DEFAULT_BASE_URL

        def log(msg, level="info"):
            if on_log:
                on_log(msg, level)

        def cancelled():
            return is_cancelled and is_cancelled()

        log("打开文档...", "info")
        self.doc = Document(self.input_path)

        log("提取文本...", "info")
        all_texts = []
        for para in self.doc.paragraphs:
            all_texts.append(para.text)
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_texts.append(cell.text)

        full_text = "\n".join(all_texts)
        if not full_text.strip():
            log("文档无文本内容", "warn")
            return

        if cancelled():
            return

        log(f"文本长度: {len(full_text)} 字符", "info")
        log(f"调用 API 按类型识别「{target}」...", "info")

        client = OpenAI(api_key=self.api_key, base_url=base_url)

        chunks = []
        step = 6000
        for start in range(0, len(full_text), step):
            chunks.append(full_text[start:start + step])

        for ci, chunk in enumerate(chunks):
            if cancelled():
                return
            if len(chunks) > 1:
                log(f"处理文本段 {ci+1}/{len(chunks)}...", "info")

            prompt = f"""你是一个精确的命名实体识别(NER)助手。请从以下文本中找出所有属于「{target}」类型的实体。

要求：
1. 返回 JSON 数组，每个元素格式: {{"text": "原文", "type": "实体类型"}}
2. text 必须是原文中的精确原文片段
3. type 是该实体的具体类别（如 "人名"、"手机号"、"身份证号"、"地址" 等）
4. 如果没有找到，返回空数组 []
5. 同一实体出现多次只需列出一次

文本内容：
{chunk}"""

            try:
                resp = client.chat.completions.create(
                    model=self.model,
                    messages=[
                        {"role": "system", "content": "你是一个精确的命名实体识别助手，只返回JSON。"},
                        {"role": "user", "content": prompt},
                    ],
                    temperature=0.1,
                    timeout=120,
                )
                raw = resp.choices[0].message.content or ""
                m = re.search(r"\[.*\]", raw, re.DOTALL)
                parsed = json.loads(m.group(0)) if m else []

                for item in parsed:
                    if isinstance(item, dict) and item.get("text"):
                        self.found_entities.append(item)
                    elif isinstance(item, str) and item:
                        self.found_entities.append({"text": item, "type": target})

            except Exception as e:
                log(f"API 调用失败: {e}", "error")

        # deduplicate
        seen = set()
        unique = []
        for ent in self.found_entities:
            key = ent["text"]
            if key not in seen:
                seen.add(key)
                unique.append(ent)
        self.found_entities = unique

        if self.found_entities:
            log(f"发现 {len(self.found_entities)} 个实体:", "found")
            for ent in self.found_entities:
                log(f"  [{ent.get('type', '?')}] {ent['text']}", "found")
        else:
            log("未发现敏感实体", "info")

        if on_names:
            on_names([f"[{e.get('type','?')}] {e['text']}" for e in self.found_entities])

        log(f"\n检测完成: {len(self.found_entities)} 个实体", "done")

    def save(self, output_path, on_log=None):
        def log(msg, level="info"):
            if on_log:
                on_log(msg, level)

        if not self.doc:
            log("无文档数据", "error")
            return 0

        entity_texts = list({e["text"] for e in self.found_entities if e.get("text")})
        if not entity_texts:
            log("无需替换", "info")
            self.doc.save(output_path)
            return 0

        text_count = 0
        mask_char = "***"
        log(f"替换 {len(entity_texts)} 类实体...", "info")

        for para in self.doc.paragraphs:
            for entity in entity_texts:
                n = _replace_in_runs(para.runs, entity, mask_char)
                text_count += n

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for entity in entity_texts:
                            n = _replace_in_runs(para.runs, entity, mask_char)
                            text_count += n

        log(f"  替换了 {text_count} 处", "found")
        log(f"保存到: {output_path}", "info")
        self.doc.save(output_path)
        log(f"保存完成: {text_count} 处替换", "done")
        return text_count
